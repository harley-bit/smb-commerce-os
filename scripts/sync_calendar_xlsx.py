#!/usr/bin/env python3
"""Sync docs/BUILD_CALENDAR.md (source of truth) into the human-facing
docs/implementation_calendar/build_calendar.xlsx dashboard.

Only touches the Status (J) and Notes (L) columns on the Daily Plan sheet.
Never touches the Card (M) / Log (N) hyperlink columns or any other sheet.

Run after closing out a card:
    python3 scripts/sync_calendar_xlsx.py
"""
import re
import sys
from pathlib import Path

import openpyxl

REPO = Path(__file__).resolve().parent.parent
MD_PATH = REPO / "docs" / "BUILD_CALENDAR.md"
XLSX_PATH = REPO / "docs" / "implementation_calendar" / "build_calendar.xlsx"

ROW_RE = re.compile(
    r"\|\s*(?:<a id=\"(?P<anchor>d\d+)\"></a>)?(?P<day>D\d+)\s*\|"
    r"\s*(?P<ph>[^|]*)\|\s*(?P<epic>[^|]*)\|\s*(?P<task>[^|]*)\|"
    r"\s*(?P<gate>[^|]*)\|\s*(?P<status>[^|]*)\|\s*(?P<card>[^|]*)\|\s*(?P<log>[^|]*)\|"
)


def parse_markdown_calendar():
    rows = {}
    text = MD_PATH.read_text()
    for line in text.splitlines():
        m = ROW_RE.match(line.strip())
        if not m:
            continue
        day = m.group("day")
        rows[day] = {
            "status": m.group("status").strip(),
            "log": m.group("log").strip(),
        }
    return rows


def main():
    if not MD_PATH.exists():
        print(f"Missing {MD_PATH}", file=sys.stderr)
        sys.exit(1)
    if not XLSX_PATH.exists():
        print(f"Missing {XLSX_PATH}", file=sys.stderr)
        sys.exit(1)

    rows = parse_markdown_calendar()
    wb = openpyxl.load_workbook(XLSX_PATH)
    ws = wb["Daily Plan"]

    updated = 0
    for row in ws.iter_rows(min_row=5):
        day_cell = row[0]
        if day_cell.value is None or not isinstance(day_cell.value, (int, float)):
            continue
        did = f"D{int(day_cell.value):03d}"
        record = rows.get(did)
        if not record:
            continue
        status_cell = ws.cell(row=day_cell.row, column=10)  # J
        note_cell = ws.cell(row=day_cell.row, column=12)    # L
        if status_cell.value != record["status"]:
            status_cell.value = record["status"]
            updated += 1
        note_text = record["log"]
        if "not yet run" in note_text.lower():
            pass  # leave Notes untouched until there's something real to say
        else:
            note_cell.value = note_text

    wb.save(XLSX_PATH)
    print(f"Synced {updated} status change(s) from {MD_PATH.name} into {XLSX_PATH}")


if __name__ == "__main__":
    main()
